# workforce/workforce/talent.py
#
# Talent search: import sourced profiles (Naukri Resdex export or any Excel/CSV),
# let the scheduler AI-score them, share a masked profile with the hiring manager,
# record the manager's answer, and invite the person to apply.
#
# Repo Python (not a Server Script) because reading Excel needs real imports.
# All data lives in the custom DocType "WF Prospect" (created by console script 11b).
#
# Privacy rules baked in here:
#   - Gender, marital status, date of birth, address, pin code, home town and
#     Naukri's activity columns are NEVER imported (not needed, must not bias scoring).
#   - Hiring managers never receive name, email, phone or the Naukri link.
#     They DO see current salary and notice period (decided Sep 11 2026).

import csv
import io
import re

import frappe
from frappe.utils import escape_html, get_url, now

HR_SENDER = "hr@cozycornerpatios.com"
MAX_ROWS = 500

RECRUITER_ROLES = ("WF HR Manager", "WF Admin", "WF Recruitment Coordinator")
HR_ROLES = ("WF HR Manager", "WF Admin")
MANAGER_ROLES = ("WF Hiring Manager", "WF CMO")

# normalised header -> WF Prospect field. Anything not listed is ignored on purpose.
HEADER_MAP = {
	"name": "full_name", "candidate name": "full_name", "full name": "full_name",
	"email id": "email", "email": "email", "email address": "email",
	"phone number": "phone", "phone": "phone", "mobile": "phone", "mobile number": "phone",
	"current location": "current_location", "location": "current_location",
	"preferred locations": "preferred_locations", "preferred location": "preferred_locations",
	"total experience": "total_experience", "experience": "total_experience",
	"curr company name": "current_company", "current company": "current_company", "company": "current_company",
	"curr company designation": "current_designation", "current designation": "current_designation", "designation": "current_designation",
	"department": "department", "role": "role", "industry": "industry",
	"key skills": "key_skills", "skills": "key_skills",
	"annual salary": "annual_salary", "current salary": "annual_salary", "current ctc": "annual_salary", "ctc": "annual_salary",
	"notice period availability to join": "notice_period", "notice period": "notice_period",
	"resume headline": "resume_headline", "headline": "resume_headline",
	"summary": "summary", "profile summary": "summary",
	"under graduation degree": "ug_degree", "ug specialization": "ug_specialization",
	"ug university institute name": "ug_institute", "ug graduation year": "ug_year",
	"post graduation degree": "pg_degree", "pg specialization": "pg_specialization",
	"pg university institute name": "pg_institute", "pg graduation year": "pg_year",
	"job title": "source_job_title", "date of application": "applied_on_source",
	"candidate profile": "profile_link", "profile link": "profile_link", "profile url": "profile_link",
}

DATA_FIELDS = (
	"full_name", "email", "phone", "current_location", "total_experience", "current_company",
	"current_designation", "department", "role", "industry", "annual_salary", "notice_period",
	"ug_degree", "ug_specialization", "ug_institute", "ug_year", "pg_degree", "pg_specialization",
	"pg_institute", "pg_year", "source_job_title", "applied_on_source",
)

# what a hiring manager may see on a shared profile
MANAGER_FIELDS = (
	"name", "job_opening", "total_experience", "current_company", "current_designation", "current_location",
	"preferred_locations", "industry", "role", "key_skills", "resume_headline", "summary",
	"annual_salary", "notice_period", "ug_degree", "ug_specialization", "ug_year",
	"pg_degree", "pg_specialization", "pg_year", "ai_score", "ai_grade", "ai_summary",
	"screening_status", "share_note", "shared_on", "shared_by", "manager_response",
	"manager_note", "responded_on", "status",
)


# ---------------------------------------------------------------- helpers
def _roles():
	return set(frappe.get_roles(frappe.session.user))


def _is_hr():
	return bool(_roles() & set(HR_ROLES))


def _require_recruiter(job_opening=None):
	roles = _roles()
	if roles & set(HR_ROLES):
		return
	if "WF Recruitment Coordinator" in roles:
		if job_opening and frappe.db.get_value("WF Job Opening", job_opening, "assigned_hr") != frappe.session.user:
			frappe.throw("You can only work on talent for positions assigned to you.")
		return
	frappe.throw("You are not permitted to use talent search.")


def _norm(h):
	return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(h or "").lower())).strip()


def _clean(v):
	if v is None:
		return ""
	s = str(v).strip()
	if s.upper() in ("NA", "N/A", "-", "NONE", "NULL", "NAN"):
		return ""
	if s.endswith(".0") and s[:-2].isdigit():
		s = s[:-2]
	return s


def _first_email(s):
	for part in re.split(r"[,;\s]+", s or ""):
		part = part.strip().lower()
		if "@" in part and "." in part.split("@")[-1]:
			return part
	return ""


def _first_phone(s):
	for part in re.split(r"[,;/]+", s or ""):
		digits = re.sub(r"[^\d+]", "", part)
		if len(re.sub(r"\D", "", digits)) >= 8:
			return digits
	return ""


def _send(recipients, subject, html, ref_name):
	"""Send from hr@ when that account exists; never let an email failure break the action."""
	try:
		sender = HR_SENDER if frappe.db.exists("Email Account", {"email_id": HR_SENDER, "enable_outgoing": 1}) else None
		frappe.sendmail(
			recipients=recipients, sender=sender, subject=subject, message=html,
			reference_doctype="WF Prospect", reference_name=ref_name,
		)
		return True
	except Exception:
		frappe.log_error(frappe.get_traceback(), "WF talent email failed: " + str(ref_name))
		return False


def _mail(title, body_html, button_label=None, button_url=None):
	btn = ""
	if button_label and button_url:
		btn = (
			'<div style="text-align:center;margin:24px 0;"><a href="' + button_url + '" '
			'style="display:inline-block;background:#4f46e5;color:#fff;padding:12px 28px;border-radius:8px;'
			'text-decoration:none;font-weight:600;">' + escape_html(button_label) + "</a></div>"
		)
	return (
		'<div style="max-width:600px;font-family:sans-serif;">'
		'<div style="background:linear-gradient(135deg,#3b3689 0%,#4f46e5 100%);padding:22px;text-align:center;">'
		'<h2 style="color:#fff;margin:0;font-size:19px;">' + escape_html(title) + "</h2></div>"
		'<div style="padding:24px;background:#fff;color:#374151;font-size:14px;line-height:1.6;">'
		+ body_html + btn +
		'<p style="color:#6b7280;font-size:12px;margin-top:20px;">Cozy Corner Patios HR Team</p></div></div>'
	)


def _job_title(job_opening):
	return frappe.db.get_value("WF Job Opening", job_opening, "job_title") or job_opening


def _default_manager(job_opening):
	req = frappe.db.get_value("WF Job Opening", job_opening, "source_requisition")
	if req:
		return frappe.db.get_value("WF Job Requisition", req, "requester") or ""
	return ""


# ---------------------------------------------------------------- file reading
def _rows_from_file(file_doc):
	"""Returns (rows, links). rows = list of lists; links = {(row_idx, col_idx): url} from cell hyperlinks."""
	fname = (file_doc.file_name or file_doc.file_url or "").lower()
	content = file_doc.get_content()
	if isinstance(content, str):
		content = content.encode("utf-8")

	if fname.endswith(".xlsx"):
		from openpyxl import load_workbook

		wb = load_workbook(io.BytesIO(content), data_only=True)  # not read_only: hyperlinks are needed
		ws = wb.worksheets[0]
		rows, links = [], {}
		for r_idx, row in enumerate(ws.iter_rows()):
			vals = []
			for c_idx, cell in enumerate(row):
				vals.append(cell.value)
				if cell.hyperlink is not None and cell.hyperlink.target:
					links[(r_idx, c_idx)] = cell.hyperlink.target
			rows.append(vals)
		return rows, links

	if fname.endswith(".xls"):
		try:
			import xlrd

			book = xlrd.open_workbook(file_contents=content)
			sh = book.sheet_by_index(0)
			return [sh.row_values(i) for i in range(sh.nrows)], {}
		except Exception:
			frappe.throw("Could not read this .xls file. Open it in Excel, save it as .xlsx, and import again.")

	if fname.endswith(".csv"):
		text = content.decode("utf-8-sig", errors="ignore")
		return [row for row in csv.reader(io.StringIO(text))], {}

	frappe.throw("Please upload an Excel (.xlsx or .xls) or CSV file.")


def _parse(rows, links):
	"""Map raw rows to WF Prospect dicts. Returns (records, header_found)."""
	header_idx = -1
	colmap = {}
	for i, row in enumerate(rows[:10]):
		cm = {}
		for c, h in enumerate(row):
			f = HEADER_MAP.get(_norm(h))
			if f and f not in cm.values():
				cm[c] = f
		if "full_name" in cm.values() and ("email" in cm.values() or "phone" in cm.values()):
			header_idx, colmap = i, cm
			break
	if header_idx < 0:
		return [], False

	link_col = next((c for c, f in colmap.items() if f == "profile_link"), None)
	records = []
	for r_off, row in enumerate(rows[header_idx + 1:]):
		r_idx = header_idx + 1 + r_off
		rec = {}
		for c, f in colmap.items():
			rec[f] = _clean(row[c]) if c < len(row) else ""
		link = links.get((r_idx, link_col)) if link_col is not None else ""
		rec["profile_link"] = link if (link or "").startswith("http") else (
			rec.get("profile_link") if (rec.get("profile_link") or "").startswith("http") else "")
		rec["email"] = _first_email(rec.get("email", ""))
		rec["phone"] = _first_phone(rec.get("phone", ""))
		for f in DATA_FIELDS:
			if rec.get(f):
				rec[f] = rec[f][:140]
		if not rec.get("full_name") and not rec.get("email"):
			continue
		records.append(rec)
	return records, True


# ---------------------------------------------------------------- APIs (HR / coordinators)
@frappe.whitelist()
def import_prospects(file_url, job_opening, source="Naukri"):
	if not job_opening or not frappe.db.exists("WF Job Opening", job_opening):
		frappe.throw("Choose the position these profiles are for.")
	_require_recruiter(job_opening)

	file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not file_name:
		frappe.throw("The uploaded file was not found. Please upload it again.")
	rows, links = _rows_from_file(frappe.get_doc("File", file_name))
	records, header_found = _parse(rows, links)
	if not header_found:
		frappe.throw("Could not find the Name and Email or Phone columns in this file.")
	if len(records) > MAX_ROWS:
		frappe.throw("This file has " + str(len(records)) + " profiles. Please import at most " + str(MAX_ROWS) + " at a time.")

	seen_emails = set(e for e in frappe.get_all("WF Prospect", filters={"job_opening": job_opening}, pluck="email") if e)
	seen_phones = set(p for p in frappe.get_all("WF Prospect", filters={"job_opening": job_opening}, pluck="phone") if p)
	applied = set((e or "").lower() for e in frappe.get_all("WF Applicant", filters={"job_opening": job_opening}, pluck="email") if e)

	batch = frappe.generate_hash(length=8)
	created, duplicates, already_applied = 0, [], []
	for rec in records:
		label = rec.get("full_name") or rec.get("email")
		if rec.get("email") and rec["email"] in applied:
			already_applied.append(label)
			continue
		if (rec.get("email") and rec["email"] in seen_emails) or (rec.get("phone") and rec["phone"] in seen_phones):
			duplicates.append(label)
			continue
		doc = frappe.new_doc("WF Prospect")
		doc.update(rec)
		doc.job_opening = job_opening
		doc.source = source if source in ("Naukri", "LinkedIn", "Referral", "Other") else "Other"
		doc.status = "New"
		doc.screening_status = "Pending"
		doc.imported_by = frappe.session.user
		doc.import_batch = batch
		doc.insert(ignore_permissions=True)
		created += 1
		if rec.get("email"):
			seen_emails.add(rec["email"])
		if rec.get("phone"):
			seen_phones.add(rec["phone"])

	frappe.db.commit()
	return {
		"created": created,
		"duplicates": duplicates[:50], "duplicate_count": len(duplicates),
		"already_applied": already_applied[:50], "already_applied_count": len(already_applied),
		"batch": batch,
		"message": str(created) + " profile(s) imported. AI scores appear within a few minutes.",
	}


@frappe.whitelist()
def get_prospects(job_opening=None):
	_require_recruiter()
	filters = {}
	if not _is_hr():
		mine = frappe.get_all("WF Job Opening", filters={"assigned_hr": frappe.session.user}, pluck="name")
		if not mine:
			return {"prospects": [], "jobs": []}
		filters["job_opening"] = ["in", mine]
	if job_opening:
		filters["job_opening"] = job_opening

	rows = frappe.get_all(
		"WF Prospect", filters=filters,
		fields=["*"], order_by="creation desc", limit_page_length=500, ignore_permissions=True,
	)
	job_names = sorted(set(r.job_opening for r in rows if r.job_opening))
	titles, managers, applicants = {}, {}, {}
	for j in job_names:
		titles[j] = _job_title(j)
		managers[j] = _default_manager(j)
		for a in frappe.get_all("WF Applicant", filters={"job_opening": j}, fields=["name", "email"], ignore_permissions=True):
			if a.email:
				applicants[(j, a.email.lower())] = a.name
	names = {}
	for r in rows:
		r["job_title"] = titles.get(r.job_opening, "")
		r["default_manager"] = managers.get(r.job_opening, "")
		r["applied_as"] = applicants.get((r.job_opening, (r.email or "").lower()), "")
		for fld in ("shared_with", "shared_by", "invited_by"):
			u = r.get(fld)
			if u and u not in names:
				names[u] = frappe.db.get_value("User", u, "full_name") or u
			r[fld + "_name"] = names.get(u, "") if u else ""
	return {"prospects": rows}


@frappe.whitelist()
def get_hiring_managers():
	_require_recruiter()
	users = {}
	for role in MANAGER_ROLES:
		for r in frappe.get_all("Has Role", filters={"role": role, "parenttype": "User"}, fields=["parent"], ignore_permissions=True):
			if r.parent in users:
				continue
			u = frappe.db.get_value("User", r.parent, ["name", "full_name", "enabled"], as_dict=True)
			if u and u.enabled:
				users[u.name] = {"email": u.name, "full_name": u.full_name or u.name}
	return sorted(users.values(), key=lambda x: x["full_name"])


@frappe.whitelist()
def share_prospect(prospect, share_with=None, note=None):
	p = frappe.get_doc("WF Prospect", prospect)
	_require_recruiter(p.job_opening)
	share_with = share_with or _default_manager(p.job_opening)
	if not share_with:
		frappe.throw("This position has no hiring manager on record. Choose who to share it with.")
	if not (set(frappe.get_roles(share_with)) & set(MANAGER_ROLES + HR_ROLES)):
		frappe.throw("You can only share with a hiring manager or the CMO.")

	p.shared_with = share_with
	p.shared_by = frappe.session.user
	p.shared_on = now()
	p.share_note = (note or "").strip()[:500]
	p.manager_response = ""
	p.manager_note = ""
	p.responded_on = None
	if p.status in ("New", "Shared", "Interested", "Not Interested"):
		p.status = "Shared"
	p.save(ignore_permissions=True)
	frappe.db.commit()

	job = _job_title(p.job_opening)
	body = (
		"<p>HR has shared a profile that could fit <strong>" + escape_html(job) + "</strong>.</p>"
		'<div style="background:#f9fafb;padding:14px 16px;border-radius:8px;margin:14px 0;">'
		"<p style=\"margin:2px 0;\"><strong>Experience:</strong> " + escape_html(p.total_experience or "-") + "</p>"
		"<p style=\"margin:2px 0;\"><strong>Currently:</strong> " + escape_html(p.current_designation or "-") + "</p>"
		"<p style=\"margin:2px 0;\"><strong>Skills:</strong> " + escape_html((p.key_skills or "-")[:200]) + "</p>"
		"</div>"
		+ ("<p><strong>Note from HR:</strong> " + escape_html(p.share_note) + "</p>" if p.share_note else "")
		+ "<p>Open it and tell HR whether you are interested.</p>"
	)
	email = frappe.db.get_value("User", share_with, "email") or share_with
	_send([email], "Profile shared with you: " + job, _mail("A profile for " + job, body, "Review the profile", get_url("/app/workforce-hub#shared")), p.name)
	return {"state": "success", "message": "Shared with " + (frappe.db.get_value("User", share_with, "full_name") or share_with) + "."}


@frappe.whitelist()
def invite_prospect(prospect):
	p = frappe.get_doc("WF Prospect", prospect)
	_require_recruiter(p.job_opening)
	if not p.email:
		frappe.throw("This profile has no email address, so no invite can be sent.")
	if frappe.db.get_value("WF Job Opening", p.job_opening, "status") != "Open":
		frappe.throw("This position is not open, so candidates can't apply to it right now.")
	if frappe.db.exists("WF Applicant", {"job_opening": p.job_opening, "email": p.email}):
		frappe.throw("This person has already applied for this position.")

	job = _job_title(p.job_opening)
	link = get_url("/apply?job_opening=" + p.job_opening)
	first = escape_html((p.full_name or "there").split(" ")[0])
	body = (
		"<p>Hi " + first + ",</p>"
		"<p>We came across your profile and think you could be a great fit for the <strong>"
		+ escape_html(job) + "</strong> role at Cozy Corner Patios.</p>"
		"<p>If you are interested, please apply using the button below. It takes about three minutes. "
		"On the form, choose <strong>" + escape_html(job) + "</strong> as the position if it isn't already selected.</p>"
	)
	sent = _send([p.email], "Invitation to apply: " + job + " at Cozy Corner Patios", _mail("An invitation to apply", body, "Apply now", link), p.name)
	if not sent:
		frappe.throw("The invite email could not be sent. Please try again later.")

	p.status = "Invited"
	p.invited_by = frappe.session.user
	p.invited_on = now()
	p.save(ignore_permissions=True)
	frappe.db.commit()
	return {"state": "success", "message": "Invite sent to " + p.email + " from hr@."}


# ---------------------------------------------------------------- APIs (hiring managers / CMO)
@frappe.whitelist()
def get_shared_prospects():
	roles = _roles()
	if not (roles & set(MANAGER_ROLES + ("WF Admin",))):
		frappe.throw("You are not permitted to view shared profiles.")
	filters = {"shared_with": frappe.session.user}
	rows = frappe.get_all(
		"WF Prospect", filters=filters, fields=list(MANAGER_FIELDS),
		order_by="shared_on desc", limit_page_length=200, ignore_permissions=True,
	)
	for r in rows:
		r["profile_ref"] = "Profile " + (r.name or "")[-5:]
		r["job_title"] = _job_title(r.job_opening)
		r["shared_by_name"] = frappe.db.get_value("User", r.shared_by, "full_name") if r.shared_by else ""
	return {"prospects": rows}


@frappe.whitelist()
def respond_prospect(prospect, response, note=None):
	if response not in ("Interested", "Not Interested"):
		frappe.throw("Choose Interested or Not interested.")
	p = frappe.get_doc("WF Prospect", prospect)
	if p.shared_with != frappe.session.user and "WF Admin" not in _roles():
		frappe.throw("This profile was not shared with you.")
	p.manager_response = response
	p.manager_note = (note or "").strip()[:500]
	p.responded_on = now()
	if p.status in ("Shared", "Interested", "Not Interested"):
		p.status = response
	p.save(ignore_permissions=True)
	frappe.db.commit()

	if p.shared_by:
		who = frappe.db.get_value("User", frappe.session.user, "full_name") or frappe.session.user
		job = _job_title(p.job_opening)
		body = (
			"<p><strong>" + escape_html(who) + "</strong> is <strong>" + escape_html(response.lower()) + "</strong> in "
			+ escape_html(p.full_name or p.name) + " for " + escape_html(job) + ".</p>"
			+ ("<p><strong>Note:</strong> " + escape_html(p.manager_note) + "</p>" if p.manager_note else "")
			+ ("<p>You can now invite them to apply.</p>" if response == "Interested" else "")
		)
		email = frappe.db.get_value("User", p.shared_by, "email") or p.shared_by
		_send([email], "Hiring manager " + response.lower() + ": " + job, _mail("Hiring manager replied", body, "Open talent search", get_url("/app/workforce-hub#talent")), p.name)
	return {"state": "success", "message": "Thanks. HR has been told."}